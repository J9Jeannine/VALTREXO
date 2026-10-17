#!/usr/bin/env python3
"""Baut die PnL-Arbeitsmappe fuer Valltrexo im Layout von "PnL KIZORA".

Ein Tab pro Monat, eine Zeile pro Tag, plus ein Tab "Settings" mit den Raten.
Die Datei wird als .xlsx geschrieben und beim Upload nach Google Sheets
konvertiert, damit die Formeln erhalten bleiben.

    python3 scripts/build_pnl_workbook.py --out /pfad/PnL_VALLTREXO.xlsx
"""
import argparse
import calendar
import datetime as dt
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EURO = '€ #,##0.00'
RATIO = '0.00'
PERCENT = '0.00%'
DATE = 'DD-MM-YYYY'

HEADER_FILL = PatternFill('solid', fgColor='1F3864')
HEADER_FONT = Font(color='FFFFFF', bold=True, size=10)
AUTO_FILL = PatternFill('solid', fgColor='E2EFDA')      # vom Skript gefuellt
MANUAL_FILL = PatternFill('solid', fgColor='FFF2CC')    # von Hand gefuellt
TOTAL_FONT = Font(bold=True)
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Spalte, Titel, Quelle. "auto" = Routine, "manual" = Hand, "formula" = berechnet.
COLUMNS = [
    ('A', 'Date', 'formula'),
    ('B', 'Revenue', 'auto'),
    ('C', 'Adspend FB', 'auto'),
    ('D', 'Adspend Pint', 'manual'),
    ('E', 'Cogs', 'formula'),
    ('F', 'SP fee', 'formula'),
    ('G', 'FB fee', 'formula'),
    ('H', 'Transaction', 'formula'),
    ('I', 'Refunds', 'auto'),
    ('J', "Abo's", 'formula'),
    ('K', 'extra Costs', 'formula'),
    ('L', 'Team Costs', 'formula'),
    ('M', 'Profit', 'formula'),
    ('N', 'ROAS', 'formula'),
    ('O', 'Profit %', 'formula'),
    ('P', 'Cogs %', 'formula'),
    ('Q', '', 'manual'),
    ('R', 'Payouts', 'formula'),
    ('S', 'Invoices', 'formula'),
    ('T', 'Extra Expenses', 'formula'),
    ('U', 'Profit after Expenses', 'formula'),
]

MONEY_COLUMNS = 'BCDEFGHIJKLMRSTU'
WIDTHS = {'A': 12, 'N': 8, 'O': 9, 'P': 9, 'Q': 3, 'U': 20}

# Settings-Tab: Zeile, Label, Wert, Hinweis. Die Zeilennummern stehen fest,
# die Monatsformeln verweisen darauf.
SETTINGS = [
    (2, 'Store', 'Valltrexo', 'dts96w-r1.myshopify.com'),
    (3, 'Currency', 'EUR', 'Shop-Waehrung'),
    (4, 'Timezone', 'Europe/Amsterdam', 'Tagesgrenze der Umsatzzahlen'),
    (5, 'Cogs %', 0.20, 'Anteil vom Revenue, bis echte Einkaufspreise stehen'),
    (6, 'SP fee %', 0.07, 'Shopify Payments'),
    (7, 'FB fee %', 0.06, 'Anteil vom Adspend'),
    (8, 'Transaction %', 0.0435, 'Zahlungsgebuehren'),
    (9, "Abo's per day", 0.0, 'Tools und Abos pro Tag'),
    (10, 'extra Costs per day', 0.0, 'sonstige laufende Kosten pro Tag'),
    (11, 'Team Costs per day', 0.0, 'Team pro Tag'),
    (12, 'Payouts per day', 0.0, 'Auszahlungen pro Tag'),
    (13, 'Invoices per day', 0.0, 'Rechnungen pro Tag'),
    (14, 'Extra Expenses per day', 0.0, 'Einmalkosten pro Tag umgelegt'),
]

SETTING_ROWS = {label: row for row, label, _, _ in SETTINGS}


def _ref(label):
    return f"Settings!$B${SETTING_ROWS[label]}"


def day_formulas(row):
    """Formeln einer Tageszeile. Schluessel ist der Spaltenbuchstabe."""
    return {
        'E': f"=ROUND(B{row}*{_ref('Cogs %')},2)",
        'F': f"=ROUND(B{row}*{_ref('SP fee %')},2)",
        'G': f"=ROUND((C{row}+D{row})*{_ref('FB fee %')},2)",
        'H': f"=ROUND(B{row}*{_ref('Transaction %')},2)",
        'J': '=' + _ref("Abo's per day"),
        'K': f"={_ref('extra Costs per day')}",
        'L': f"={_ref('Team Costs per day')}",
        'M': (f"=B{row}-C{row}-D{row}-E{row}-F{row}-G{row}"
              f"-H{row}-I{row}-J{row}-K{row}-L{row}"),
        'N': f"=IF((C{row}+D{row})=0,0,B{row}/(C{row}+D{row}))",
        'O': f"=IF(B{row}=0,0,M{row}/B{row})",
        'P': f"=IF(B{row}=0,0,E{row}/B{row})",
        'R': f"={_ref('Payouts per day')}",
        'S': f"={_ref('Invoices per day')}",
        'T': f"={_ref('Extra Expenses per day')}",
        'U': f"=M{row}-R{row}-S{row}-T{row}",
    }


def build_settings(sheet):
    sheet['A1'] = 'Setting'
    sheet['B1'] = 'Value'
    sheet['C1'] = 'Note'
    for cell in ('A1', 'B1', 'C1'):
        sheet[cell].fill = HEADER_FILL
        sheet[cell].font = HEADER_FONT

    for row, label, value, note in SETTINGS:
        sheet[f'A{row}'] = label
        sheet[f'B{row}'] = value
        sheet[f'C{row}'] = note
        sheet[f'B{row}'].fill = MANUAL_FILL
        sheet[f'B{row}'].border = BORDER
        if label.endswith('%'):
            sheet[f'B{row}'].number_format = PERCENT
        elif 'per day' in label:
            sheet[f'B{row}'].number_format = EURO

    note_row = len(SETTINGS) + 3
    sheet[f'A{note_row}'] = 'Legende'
    sheet[f'A{note_row}'].font = TOTAL_FONT
    sheet[f'A{note_row + 1}'] = 'gruen'
    sheet[f'B{note_row + 1}'] = 'traegt die Routine ein (Revenue, Adspend FB, Refunds)'
    sheet[f'A{note_row + 1}'].fill = AUTO_FILL
    sheet[f'A{note_row + 2}'] = 'gelb'
    sheet[f'B{note_row + 2}'] = 'von Hand eintragen'
    sheet[f'A{note_row + 2}'].fill = MANUAL_FILL
    sheet[f'A{note_row + 3}'] = 'weiss'
    sheet[f'B{note_row + 3}'] = 'Formel, nicht ueberschreiben'

    sheet.column_dimensions['A'].width = 24
    sheet.column_dimensions['B'].width = 24
    sheet.column_dimensions['C'].width = 52


def build_month(sheet, year, month, seed=None):
    """seed: Dict Datum (YYYY-MM-DD) -> {'revenue': float, 'refunds': float}."""
    seed = seed or {}
    days = calendar.monthrange(year, month)[1]

    for letter, title, source in COLUMNS:
        cell = sheet[f'{letter}1']
        cell.value = title
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sheet.column_dimensions[letter].width = WIDTHS.get(letter, 13)

    first_row = 2
    last_row = first_row + days - 1

    for index in range(days):
        row = first_row + index
        formulas = day_formulas(row)
        date = dt.date(year, month, index + 1)
        seeded = seed.get(date.isoformat(), {})

        date_cell = sheet[f'A{row}']
        date_cell.value = date
        date_cell.number_format = DATE

        for letter, _, source in COLUMNS:
            cell = sheet[f'{letter}{row}']
            cell.border = BORDER
            if letter in formulas:
                cell.value = formulas[letter]
            elif source == 'auto':
                cell.fill = AUTO_FILL
                if letter == 'B' and 'revenue' in seeded:
                    cell.value = seeded['revenue']
                elif letter == 'I' and 'refunds' in seeded:
                    cell.value = seeded['refunds']
            elif source == 'manual' and letter != 'Q':
                cell.fill = MANUAL_FILL
            if letter in MONEY_COLUMNS:
                cell.number_format = EURO
            elif letter == 'N':
                cell.number_format = RATIO
            elif letter in ('O', 'P'):
                cell.number_format = PERCENT

    total_row = last_row + 1
    sheet[f'A{total_row}'] = 'TOTAL'
    sheet[f'A{total_row}'].font = TOTAL_FONT

    for letter, _, _ in COLUMNS:
        if letter in ('A', 'Q'):
            continue
        cell = sheet[f'{letter}{total_row}']
        cell.font = TOTAL_FONT
        cell.border = BORDER
        if letter == 'N':
            cell.value = (f"=IF((C{total_row}+D{total_row})=0,0,"
                          f"B{total_row}/(C{total_row}+D{total_row}))")
            cell.number_format = RATIO
        elif letter == 'O':
            cell.value = f"=IF(B{total_row}=0,0,M{total_row}/B{total_row})"
            cell.number_format = PERCENT
        elif letter == 'P':
            cell.value = f"=IF(B{total_row}=0,0,E{total_row}/B{total_row})"
            cell.number_format = PERCENT
        else:
            cell.value = f"=SUM({letter}{first_row}:{letter}{last_row})"
            cell.number_format = EURO

    sheet.freeze_panes = 'B2'


def month_tab_name(year, month):
    return f"{calendar.month_name[month]} {year % 100:02d}"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--out', required=True)
    parser.add_argument('--start', default='2026-06', help='erster Monat, YYYY-MM')
    parser.add_argument('--months', type=int, default=6)
    parser.add_argument('--seed', help='JSON von pnl_revenue.py --json, fuellt Revenue und Refunds')
    args = parser.parse_args()

    seed = {}
    if args.seed:
        with open(args.seed) as handle:
            for row in json.load(handle):
                seed[row['date']] = row

    year, month = (int(part) for part in args.start.split('-'))

    workbook = Workbook()
    build_settings(workbook.active)
    workbook.active.title = 'Settings'

    for _ in range(args.months):
        build_month(workbook.create_sheet(month_tab_name(year, month)), year, month, seed)
        month += 1
        if month == 13:
            year, month = year + 1, 1

    workbook.save(args.out)
    print(f"wrote {args.out} with tabs: {', '.join(workbook.sheetnames)}")


if __name__ == '__main__':
    main()
